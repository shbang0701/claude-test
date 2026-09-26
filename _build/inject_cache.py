# -*- coding: utf-8 -*-
"""수식 셀에 계산된 값(<v>)을 미리 넣는다.
openpyxl은 값 없이 수식만 저장하므로, LibreOffice로 사본을 재계산한 뒤 그 값을 원본 XML에 옮긴다.
(원본의 서식·인쇄 설정은 그대로 두고 값만 추가 — Excel은 열 때 다시 계산한다.)"""
import os
import shutil
import sys
import tempfile
import zipfile

import openpyxl
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lo  # noqa: E402

ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def sheet_paths(z):
    wb = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rmap = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for s in wb.find("m:sheets", NS):
        rid = s.get("{%s}id" % NS["r"])
        t = rmap[rid]
        out[s.get("name")] = "xl/" + t.lstrip("/").replace("xl/", "") if not t.startswith("xl/") else t
    return out


def inject(path):
    tmpd = tempfile.mkdtemp()
    cp = lo.convert(path, "xlsx", tmpd)          # 변환하면서 수식이 계산된 사본
    if not os.path.exists(cp):
        raise SystemExit("LibreOffice 변환 실패: " + path)
    vals = openpyxl.load_workbook(cp, data_only=True)
    bad = [(ws.title, c.coordinate, c.value) for ws in vals.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value in ERRORS]
    if bad:
        raise SystemExit(f"수식 오류 {len(bad)}건: {bad[:10]}")
    tmp = path + ".tmp"
    n = 0
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        paths = sheet_paths(zin)
        inv = {v: k for k, v in paths.items()}
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in inv:
                ws = vals[inv[item.filename]]
                root = etree.fromstring(data)
                for c in root.iter("{%s}c" % NS["m"]):
                    f = c.find("m:f", NS)
                    if f is None:
                        continue
                    v = ws[c.get("r")].value
                    ve = c.find("m:v", NS)
                    if ve is None:
                        ve = etree.SubElement(c, "{%s}v" % NS["m"])
                    if v is None:
                        v = ""
                    if isinstance(v, bool):
                        c.set("t", "b")
                        ve.text = "1" if v else "0"
                    elif isinstance(v, (int, float)):
                        if "t" in c.attrib:
                            del c.attrib["t"]
                        ve.text = repr(v) if isinstance(v, float) else str(v)
                    else:
                        c.set("t", "str")
                        ve.text = str(v)
                    n += 1
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    shutil.move(tmp, path)
    shutil.rmtree(tmpd, ignore_errors=True)
    return n


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(os.path.basename(p), inject(p), "formula values injected")
