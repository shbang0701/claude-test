# -*- coding: utf-8 -*-
"""구성 요소 슬라이드: 장비 · 랙 · 영역 · 연결선 · 라벨 · 범례 · 표 · 비교 · 현황 · 절차."""
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Pt
from ppt_kit import (P, I, text, shape, rect, rrect, oval, line, connect, group, icon, slide, table, rgb, set_font,
                     line_fmt)
from ppt_elements import (LINES, LSTYLE, ZONES, KIND_TINT, node_card, node_compact, node_tile, node_box, pill, badge,
                          tag, zone, line_sample, line_legend, color_legend, kpi, usage_bar, steps, checkbox, rack)


def cap(sl, x, y, s, w=6):
    return text(sl, x, y, w, 0.22, s, 9, True, P["s500"], name="설명 캡션")


# ── 장비 요소 ────────────────────────────────────────────────
def nodes(prs):
    sl = slide(prs, "장비 요소", "구성 요소 · 장비")
    cap(sl, 0.5, 1.3, "카드형 — 기본 (흰 바탕 · 옅은 바탕 · 강조)")
    node_card(sl, 0.5, 1.6, "server", "WEB-01", "10.10.1.11 · 웹 서버", P["svc"])
    node_card(sl, 2.8, 1.6, "database", "DB-01", "10.10.2.21 · Oracle", P["san"], style="tint")
    node_card(sl, 5.1, 1.6, "firewall", "FW-01", "외부 방화벽 · Active", P["crit"], style="solid")
    cap(sl, 7.6, 1.3, "타일형 — 아이콘을 색 사각형에")
    node_tile(sl, 7.6, 1.62, "switch", "CORE-SW", "L3 · 100G", P["mgmt"], w=1.75)
    node_tile(sl, 9.45, 1.62, "storage", "STG-01", "Unity XT 480", P["san"], w=1.75)
    node_tile(sl, 11.3, 1.62, "cloud", "AWS", "VPC 10.50.0.0/16", P["svc"], w=1.6)
    cap(sl, 0.5, 2.55, "상자형 — 정보 여러 줄")
    node_box(sl, 0.5, 2.85, "server", "APP-01", ["10.10.1.31", "RHEL 8 · 8C/32G"], P["svc"])
    node_box(sl, 2.2, 2.85, "vm", "VM-APP-02", ["10.10.1.32", "vSphere 7"], P["ic"])
    node_box(sl, 3.9, 2.85, "load-balancer", "L4-01", ["VIP 10.10.1.10", "Active"], P["mgmt"])
    cap(sl, 5.85, 2.55, "간단형 — 아이콘 + 이름 (빽빽한 구성도)")
    for k, (key, nm, sub, col) in enumerate([("server", "서버", "x 12", P["svc"]), ("switch", "스위치", "L2", P["mgmt"]),
                                             ("firewall", "방화벽", None, P["crit"]), ("storage", "스토리지", None, P["san"]),
                                             ("database", "DB", None, P["san"]), ("cloud", "클라우드", None, P["svc"]),
                                             ("user", "사용자", None, P["s600"])]):
        node_compact(sl, 6.0 + k * 0.98, 2.9, key, nm, sub, 0.5, col, w=0.95)
    cap(sl, 0.5, 4.2, "이중화 묶음")
    g1, b1 = node_card(sl, 0.5, 4.52, "switch", "TOR-A", "Nexus 93180 · U39", P["mgmt"], w=2.2, h=0.56)
    g2, b2 = node_card(sl, 0.5, 5.14, "switch", "TOR-B", "Nexus 93180 · U38", P["mgmt"], w=2.2, h=0.56)
    connect(sl, b1, 3, b2, 3, "elbow", LSTYLE["ha"][0], LSTYLE["ha"][1], LSTYLE["ha"][2], name="연결-피어링크", bulge=0.25)
    pill(sl, 2.85, 4.98, "vPC", P["ic"])
    tag(sl, 1.9, 4.28, "Active / Active", "info", w=1.0)
    cap(sl, 4.0, 4.2, "클러스터 · 그룹 (점선 = 논리 묶음)")
    zone(sl, 4.0, 4.5, 4.1, 1.55, "group", "ESXi 클러스터 · 3 Nodes")
    for k in range(3):
        node_compact(sl, 4.45 + k * 1.3, 4.85, "virtual-host", f"ESX-0{k+1}", "10.10.3.1" + str(k + 1), 0.46, P["ic"], w=1.2)
    cap(sl, 8.5, 4.2, "변경 표시 (신규 · 철거)")
    g, b = node_card(sl, 8.5, 4.52, "server", "APP-03", "신규 도입", P["ok"])
    line_fmt(b, P["ok"], 1.25, "dash")
    tag(sl, 10.2, 4.4, "신규", "new", solid=True, w=0.45, h=0.22, size=8)
    g, b = node_card(sl, 8.5, 5.3, "server", "OLD-APP", "철거 예정", P["s400"], icon_color=P["s400"])
    line_fmt(b, P["crit"], 1.0, "sysDash")
    tag(sl, 10.2, 5.18, "철거", "del", solid=True, w=0.45, h=0.22, size=8)
    text(sl, 0.5, 6.35, 12.3, 0.5, [[("사용법  ", {"bold": True, "color": P["svc"]}),
                                     ("모든 장비 요소는 그룹입니다. 연결선은 '바탕' 도형(그룹 안 사각형)에 붙이면 요소를 옮겨도 따라옵니다. "
                                      "아이콘 교체: 아이콘만 선택 → 삭제 → 아이콘 쪽에서 새 아이콘을 복사해 같은 자리에 붙여넣기.", {})]],
         9, False, P["s700"])
    return sl


# ── 랙 ─────────────────────────────────────────────────────
def racks(prs):
    sl = slide(prs, "랙 · 하드웨어 요소", "구성 요소 · 랙")
    devs1 = [(42, 1, "PP-01 패치패널", "기타"), (41, 1, "MGMT-SW", "네트워크"), (39, 1, "TOR-A", "네트워크"),
             (38, 1, "TOR-B", "네트워크"), (36, 1, "SAN-A", "스토리지"), (35, 1, "SAN-B", "스토리지"),
             (30, 2, "STG-01  Unity XT", "스토리지"), (20, 2, "SRV-01  DL380", "서버"), (18, 2, "SRV-02  DL380", "서버"),
             (16, 1, "SRV-03  R640", "서버"), (1, 10, "ENC-01  c7000", "서버")]
    rack(sl, 0.9, 1.65, devices=devs1, title_txt="DC1-R02 (전면)", gname="랙-DC1-R02")
    devs2 = [(40, 2, "FW-01", "보안"), (38, 2, "FW-02", "보안"), (35, 1, "L4-01", "네트워크"), (34, 1, "L4-02", "네트워크"),
             (30, 2, "UPS 관리", "전원"), (20, 4, "BAK-01  백업", "기타")]
    rack(sl, 3.85, 1.65, devices=devs2, title_txt="DC1-R03 (전면)", gname="랙-DC1-R03")
    x = 7.0
    cap(sl, x, 1.3, "U 블록 (랙 안에 복사해서 쌓기)")
    uh = 0.118
    for k, (h, nm, kind) in enumerate([(1, "1U 장비", "네트워크"), (2, "2U 장비", "서버"), (4, "4U 장비", "스토리지"),
                                       (1, "빈 1U (블랭크)", "빈칸")]):
        yy = 1.65 + k * 0.62
        rect(sl, x, yy, 2.0, h * uh, fill=KIND_TINT.get(kind) or P["white"], line=P["s600"] if kind != "빈칸" else P["s300"],
             lw=0.5, dash=None if kind != "빈칸" else "sysDash", txt=nm, size=6.5 if h == 1 else 8, bold=True,
             color=P["ink"] if kind != "빈칸" else P["s400"], align="l", margins=(0.06, 0.02, 0, 0), name=f"U블록-{nm}")
        text(sl, x + 2.1, yy, 1.6, 0.2, f"높이 {h} × 0.118in", 8, False, P["s500"])
    color_legend(sl, x, 4.3, [("서버", KIND_TINT["서버"]), ("네트워크", KIND_TINT["네트워크"]),
                              ("스토리지", KIND_TINT["스토리지"]), ("보안", KIND_TINT["보안"]),
                              ("전원", KIND_TINT["전원"]), ("기타", KIND_TINT["기타"])],
                 title_txt="구분 색 (Excel 랙 시트와 같음)", colw=0.95)
    text(sl, x, 5.25, 5.8, 1.3, [
        [("축척", {"bold": True}), ("  1U = 0.118in (42U 랙 = 약 5in). 블록은 U 수 × 0.118in 높이로 만들어 두었습니다.", {})],
        [("쌓기", {"bold": True}), ("  블록을 복사해 랙 안에 놓고, [정렬 > 왼쪽 맞춤]으로 줄을 맞춥니다.", {})],
        [("Excel 연계", {"bold": True}), ("  정확한 실장도는 Excel 랙 파일의 [랙] 시트(자동 생성)를 '그림으로 복사'해 붙여도 됩니다.", {})]],
         9.5, False, P["s700"], spacing=4, line_spacing=1.1)
    return sl


# ── 영역 ────────────────────────────────────────────────────
def zones(prs):
    sl = slide(prs, "영역 박스", "구성 요소 · 영역")
    cap(sl, 0.5, 1.3, "망 · 논리 영역 — 옅은 바탕 + 용도 색 테두리 (연결선 색과 같은 규칙)")
    order = ["ext", "dmz", "svc", "mgmt", "bak", "san", "dr", "cloud"]
    for k, kind in enumerate(order):
        x = 0.5 + (k % 4) * 3.1
        y = 1.65 + (k // 4) * 1.55
        zone(sl, x, y, 2.9, 1.35, kind, icon_key={"cloud": "cloud", "ext": "internet"}.get(kind))
    cap(sl, 0.5, 4.85, "물리 영역 · 그룹 · 보안 경계")
    zone(sl, 0.5, 5.3, 3.9, 1.45, "phys", "DC1 · 2F 전산실", tab=True)
    zone(sl, 4.7, 5.3, 3.6, 1.45, "group", "업무 그룹 (논리)")
    zone(sl, 8.6, 5.3, 4.2, 1.45, "sec", "보안 경계 (방화벽 안쪽)")
    text(sl, 0.7, 5.72, 3.5, 0.9, ["이름표형(탭): 센터 · 층 · 전산실처럼", "큰 물리 영역에 사용"], 9, False, P["s500"],
         line_spacing=1.2)
    text(sl, 4.9, 5.72, 3.2, 0.9, ["점선 = 실제 경계가 아닌 논리 묶음", "(클러스터, 업무 그룹)"], 9, False, P["s500"],
         line_spacing=1.2)
    text(sl, 8.8, 5.72, 3.8, 0.9, ["굵은 빨간 점선 = 보안 경계", "영역은 맨 뒤로 보내기 (Ctrl+Shift+[ )"], 9, False, P["s500"],
         line_spacing=1.2)
    return sl


# ── 연결선 ───────────────────────────────────────────────────
def lines_slide(prs):
    sl = slide(prs, "연결선 · 선 규칙", "구성 요소 · 연결선")
    line_legend(sl, 0.5, 1.3, cols=1, colw=5.0, rowh=0.44, title_txt="연결 종류 (색 = 용도, 모양 = 구분)")
    text(sl, 0.5, 5.2, 5.3, 1.5, [
        [("규칙", {"bold": True, "color": P["svc"]})],
        "· 선 색은 Excel 도면의 포트 색과 같습니다 (서비스 = 파랑, SAN = 주황 …)",
        "· 꺾임은 직각(꺾인 연결선), 교차점에는 점을 찍지 않습니다",
        "· 속도 · 회선 수는 선 위 라벨(알약)로, 방향이 중요할 때만 화살표"], 9.5, False, P["s700"], line_spacing=1.15)
    x0 = 6.3
    cap(sl, x0, 1.3, "① 연결점에 붙인 꺾인 연결선 + 속도 라벨")
    g1, a = node_card(sl, x0, 1.62, "server", "SRV-01", "DL380 Gen10", P["svc"], w=1.9, h=0.56)
    g2, b = node_card(sl, x0 + 4.2, 1.62, "switch", "TOR-A", "Eth1/11", P["mgmt"], w=1.9, h=0.56)
    connect(sl, a, 3, b, 1, "elbow", *LSTYLE["svc"], name="연결-서비스")
    pill(sl, x0 + 2.55, 1.79, "25G ×2", P["svc"])
    cap(sl, x0, 2.5, "② 이중화 A/B (두 스위치로 나눠 연결)")
    g3, s = node_card(sl, x0, 2.95, "server", "SRV-02", "", P["svc"], w=1.9, h=0.56)
    g4, ta = node_card(sl, x0 + 4.2, 2.8, "switch", "SAN-A", "", P["san"], w=1.9, h=0.42)
    g5, tb = node_card(sl, x0 + 4.2, 3.35, "switch", "SAN-B", "", P["san"], w=1.9, h=0.42)
    connect(sl, s, 3, ta, 1, "elbow", *LSTYLE["san"], name="연결-SAN A")
    connect(sl, s, 3, tb, 1, "elbow", *LSTYLE["san"], name="연결-SAN B")
    pill(sl, x0 + 2.35, 2.8, "Fabric A", P["san"])
    pill(sl, x0 + 2.35, 3.52, "Fabric B", P["san"])
    cap(sl, x0, 4.05, "③ 묶음(LACP) · ④ 방향 · ⑤ 관리 · 전원")
    y = 4.72
    line(sl, x0, y, x0 + 2.2, y, P["svc"], 1.75)
    line(sl, x0, y + 0.14, x0 + 2.2, y + 0.14, P["svc"], 1.75)
    oval(sl, x0 + 0.95, y - 0.1, 0.3, 0.34, fill=None, line=P["s700"], lw=1.0, name="묶음 표시")
    pill(sl, x0 + 1.3, y - 0.34, "LACP", P["s700"], w=0.55)
    line(sl, x0 + 2.8, y + 0.07, x0 + 4.9, y + 0.07, P["svc"], 1.75, tail="triangle", name="방향선")
    text(sl, x0 + 2.8, y + 0.15, 2.2, 0.2, "데이터 흐름 방향 (필요할 때만)", 8, False, P["s500"])
    line(sl, x0 + 5.3, y, x0 + 6.5, y, *LSTYLE["mgmt"], name="관리선")
    line(sl, x0 + 5.3, y + 0.2, x0 + 6.5, y + 0.2, *LSTYLE["pwr"], name="전원선")
    text(sl, x0 + 5.3, y + 0.28, 1.4, 0.2, "관리 · 전원", 8, False, P["s500"])
    cap(sl, x0, 5.15, "⑥ 선 끝 붙이기")
    text(sl, x0, 5.45, 6.4, 1.3, [
        "· 선 끝을 도형 가장자리의 연결점(점)에 놓으면 붙습니다 → 도형을 옮겨도 선이 따라옴",
        "· 꺾이는 위치는 선 가운데 노란 조절점을 끌어서 옮깁니다",
        "· 선 서식만 복사: 선 선택 → 서식 복사(Ctrl+Shift+C) → 다른 선 선택 → Ctrl+Shift+V"],
         9.5, False, P["s700"], line_spacing=1.2)
    return sl


# ── 라벨·태그 ────────────────────────────────────────────────
def labels(prs):
    sl = slide(prs, "라벨 · 태그 · 번호", "구성 요소 · 라벨")
    cap(sl, 0.5, 1.3, "포트 · 속도 · 주소 라벨 (알약)")
    x = 0.5
    for t, c in (("Eth1/1", P["svc"]), ("Gi1/0/12", P["mgmt"]), ("P0", P["san"]), ("25G ×2", P["svc"]),
                 ("100G", P["wan"]), ("VLAN 100", P["svc"]), ("10.10.1.0/24", P["s700"]), ("iLO", P["mgmt"])):
        p = pill(sl, x, 1.62, t, c)
        x += p.width / 914400 + 0.12
    cap(sl, 0.5, 2.1, "주소 · VLAN 태그 (사각)")
    x = 0.5
    for t, kind in (("VLAN 10  서비스", "info"), ("VLAN 20  관리", "ok"), ("VLAN 30  백업", "mute"), ("IP 10.10.1.11", "mute")):
        tg = tag(sl, x, 2.42, t, kind)
        x += tg.width / 914400 + 0.12
    cap(sl, 0.5, 2.9, "번호 배지 — 도면 · 절차 · 주석 번호")
    for k in range(1, 10):
        badge(sl, 0.5 + (k - 1) * 0.36, 3.22, k)
    badge(sl, 3.8, 3.22, "A", fill=P["svc"])
    badge(sl, 4.16, 3.22, "B", fill=P["san"])
    text(sl, 4.6, 3.22, 3.2, 0.26, "Excel 도면의 '도면 번호'와 같은 번호를 쓰면 좋습니다", 8.5, False, P["s500"], anchor="m")
    cap(sl, 0.5, 3.75, "상태 태그")
    x = 0.5
    for t, kind in (("● 정상", "ok"), ("● 주의", "warn"), ("● 장애", "crit"), ("● 점검 중", "info"), ("● 예정", "mute")):
        tg = tag(sl, x, 4.07, t, kind)
        x += tg.width / 914400 + 0.12
    x = 5.0
    for t, kind in (("정상", "ok"), ("주의", "warn"), ("장애", "crit"), ("점검 중", "info")):
        tg = tag(sl, x, 4.07, t, kind, solid=True)
        x += tg.width / 914400 + 0.12
    cap(sl, 0.5, 4.6, "변경 태그 — 작업 · 변경 계획서")
    x = 0.5
    for t, kind in (("신규", "new"), ("변경", "chg"), ("철거", "del"), ("유지", "keep")):
        tg = tag(sl, x, 4.92, t, kind, solid=kind != "keep", w=0.62)
        x += 0.74
    text(sl, 3.5, 4.92, 5, 0.24, "신규 = 초록 · 변경 = 주황 · 철거 = 빨강 · 유지 = 회색 (장비 테두리 점선과 함께)", 8.5, False,
         P["s500"], anchor="m")
    # 오른쪽: 사용 예
    x0 = 8.4
    cap(sl, x0, 1.3, "사용 예")
    zone(sl, x0, 1.65, 4.4, 3.9, "svc", "서비스망 VLAN 10")
    g1, a = node_card(sl, x0 + 0.3, 2.15, "server", "SRV-01", "10.10.1.21", P["svc"], w=1.8, h=0.56)
    g2, b = node_card(sl, x0 + 0.3, 3.6, "switch", "TOR-A", "10.10.0.39", P["mgmt"], w=1.8, h=0.56)
    connect(sl, a, 2, b, 0, "elbow", *LSTYLE["svc"])
    pill(sl, x0 + 1.25, 3.0, "Eth1/11", P["svc"])
    badge(sl, x0 + 2.2, 2.1, 1, d=0.24, size=9)
    badge(sl, x0 + 2.2, 3.55, 2, d=0.24, size=9)
    tag(sl, x0 + 2.6, 2.3, "● 정상", "ok")
    tag(sl, x0 + 2.6, 3.75, "변경", "chg", solid=True, w=0.55)
    text(sl, x0 + 0.3, 4.55, 3.9, 0.8, ["① SRV-01 : 25G 이중화 구성 완료", "② TOR-A : 포트 11 VLAN 변경 (10 → 110)"],
         9, False, P["s700"], line_spacing=1.2)
    text(sl, 0.5, 5.6, 7.5, 1.0, [[("팁  ", {"bold": True, "color": P["svc"]}),
                                   ("알약·태그는 글자를 바꾸면 폭이 자동으로 늘지 않습니다. 오른쪽 조절점으로 폭을 맞추거나 "
                                    "[도형 서식 > 텍스트 상자 > 도형을 텍스트 크기에 맞춤]을 켜세요.", {})]],
         9, False, P["s700"], line_spacing=1.15)
    return sl


# ── 범례·주석·강조 ───────────────────────────────────────────
def legends(prs):
    sl = slide(prs, "범례 · 주석 · 강조", "구성 요소 · 범례")
    cap(sl, 0.5, 1.3, "범례 — 연결선 (2열 요약형)")
    lg = line_legend(sl, 0.5, 1.6, keys=["svc", "san", "mgmt", "bak", "ha", "wan"], cols=2, colw=2.2, rowh=0.3, desc=False,
                     title_txt=None, gname="범례-연결선 요약")
    cap(sl, 0.5, 2.7, "범례 — 장비 구분 색")
    color_legend(sl, 0.5, 2.95, [("서버", KIND_TINT["서버"]), ("네트워크", KIND_TINT["네트워크"]), ("스토리지", KIND_TINT["스토리지"]),
                                 ("보안", KIND_TINT["보안"])], title_txt=None, colw=1.05)
    cap(sl, 0.5, 3.85, "범례 — 아이콘")
    items = [("server", "서버"), ("switch", "스위치"), ("firewall", "방화벽"), ("storage", "스토리지"), ("cloud", "클라우드")]
    g = []
    for k, (key, nm) in enumerate(items):
        g.append(icon(sl, key, 0.5 + k * 0.92, 4.15, 0.3, P["navy"]))
        g.append(text(sl, 0.85 + k * 0.92, 4.15, 0.6, 0.3, nm, 8.5, False, P["ink"], anchor="m", wrap=False))
    group(sl, g, "범례-아이콘")
    cap(sl, 0.5, 4.75, "각주 · 출처")
    text(sl, 0.5, 5.05, 5.5, 0.5, ["※ 2026.09 기준 구성. IP는 관리망 기준.", "출처: 인프라팀 자산 관리 대장 (2026-09-26)"],
         8, False, P["s500"], line_spacing=1.2, name="각주")
    x0 = 6.4
    cap(sl, x0, 1.3, "말풍선 (설명 달기)")
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, x0, 1.6, 2.9, 0.75, fill=P["white"], line=P["svc"], lw=1.0,
          txt=[[("변경 포인트  ", {"bold": True, "color": P["svc"]}), ("업링크 40G → 100G", {})]], size=9.5,
          align="l", margins=(0.1, 0.08, 0.05, 0.05), adj={0: -0.2, 1: 0.95}, name="말풍선")
    cap(sl, x0 + 3.4, 1.3, "주석 상자 · 경고 상자")
    g = [rrect(sl, x0 + 3.4, 1.6, 3.0, 0.62, 0.05, fill=P["t_info"], line=None, name="바탕"),
         icon(sl, "info", x0 + 3.52, 1.75, 0.3, P["info"]),
         text(sl, x0 + 3.95, 1.66, 2.35, 0.5, ["참고", "작업 중 관리망 접속 불가 (약 5분)"], 8.5, False, P["ink"], line_spacing=1.1)]
    group(sl, g, "주석-참고")
    g = [rrect(sl, x0 + 3.4, 2.35, 3.0, 0.62, 0.05, fill=P["t_warn"], line=None, name="바탕"),
         icon(sl, "warning", x0 + 3.52, 2.5, 0.3, P["warn"]),
         text(sl, x0 + 3.95, 2.41, 2.35, 0.5, ["주의", "PSU 교체 시 A/B 순서 지킬 것"], 8.5, False, P["ink"], line_spacing=1.1)]
    group(sl, g, "주석-주의")
    cap(sl, x0, 3.2, "강조 테두리 (작업 · 변경 구간)")
    zone(sl, x0, 3.55, 6.4, 1.9, "svc", "서비스망")
    node_card(sl, x0 + 0.3, 4.05, "server", "SRV-01", "", P["svc"], w=1.6, h=0.5)
    node_card(sl, x0 + 2.3, 4.05, "server", "SRV-02", "", P["svc"], w=1.6, h=0.5)
    node_card(sl, x0 + 4.3, 4.05, "server", "SRV-03", "", P["svc"], w=1.6, h=0.5)
    g = [rrect(sl, x0 + 2.15, 3.85, 3.95, 1.3, 0.08, fill=None, line=P["crit"], lw=1.75, dash="dash", name="강조 테두리"),
         tag(sl, x0 + 4.9, 3.72, "작업 구간", "crit", solid=True, w=0.95, h=0.24, size=8.5)]
    group(sl, g, "강조-작업 구간")
    cap(sl, x0, 5.65, "도면 번호 설명 (Excel 장비 도면과 같은 표기)")
    g = [badge(sl, x0, 5.95, 7, d=0.26), text(sl, x0 + 0.35, 5.95, 1.4, 0.26, "도면 번호", 9, False, P["ink"], anchor="m"),
         text(sl, x0 + 1.6, 5.95, 0.5, 0.26, "LOM1", 8, False, P["s500"], anchor="m"),
         text(sl, x0 + 2.1, 5.95, 1.5, 0.26, "포트명(장비 표기)", 9, False, P["ink"], anchor="m"),
         rrect(sl, x0 + 3.75, 5.96, 0.95, 0.24, 0.02, fill="FFF1BF", line=None, txt="R02-D001", size=8, color=P["ink"],
               name="라벨칸"),
         text(sl, x0 + 4.8, 5.95, 1.3, 0.26, "케이블 라벨", 9, False, P["ink"], anchor="m")]
    group(sl, g, "범례-표기 구분")
    return sl


# ── 표 ─────────────────────────────────────────────────────
def tables(prs):
    sl = slide(prs, "표", "구성 요소 · 표")
    cap(sl, 0.5, 1.3, "장비 목록")
    rows = [["장비명", "구분", "모델", "위치", "관리 IP", "용도"],
            ["SRV-01", "서버", "DL380 Gen10", "R02 U20", "10.10.1.21", "WEB"],
            ["SRV-02", "서버", "DL380 Gen10", "R02 U18", "10.10.1.22", "WAS"],
            ["TOR-A", "네트워크", "Nexus 93180YC", "R02 U39", "10.10.0.39", "ToR"],
            ["STG-01", "스토리지", "Unity XT 480", "R02 U30", "10.10.0.30", "공유 스토리지"]]
    table(sl, 0.5, 1.6, [1.0, 0.85, 1.35, 0.85, 1.05, 1.1], rows, row_h=0.3, name="표-장비 목록")
    cap(sl, 7.0, 1.3, "IP · VLAN 할당")
    rows = [["VLAN", "이름", "대역", "게이트웨이", "비고"],
            ["10", "서비스", "10.10.1.0/24", "10.10.1.1", "HSRP"],
            ["20", "관리", "10.10.0.0/24", "10.10.0.1", "OOB"],
            ["30", "백업", "10.10.5.0/24", "10.10.5.1", ""],
            ["40", "vMotion", "10.10.6.0/24", "—", "L2 전용"]]
    fills = {(r, 0): c for r, c in ((1, P["t_svc"]), (2, P["t_mgmt"]), (3, P["t_bak"]), (4, P["t_ic"]))}
    table(sl, 7.0, 1.6, [0.65, 0.95, 1.35, 1.15, 1.7], rows, row_h=0.3, name="표-VLAN", align=["c", "l", "l", "l", "l"],
          fills=fills)
    cap(sl, 0.5, 3.35, "포트 연결표 (Excel 표와 같은 열 순서)")
    rows = [["번호", "포트명", "용도", "상대 장비", "상대 포트", "케이블 라벨"],
            ["1", "S1-P1", "스토리지", "SAN-A", "P0", "R02-D038"],
            ["5", "FLR1", "서비스", "TOR-A", "Eth1/11", "R02-D013"],
            ["7", "iLO", "관리", "MGMT-SW", "Gi1/0/1", "R02-D001"],
            ["10", "PSU1", "전원", "PDU-A", "C13-7", "R02-P013"]]
    fills = {(1, 0): P["san"], (2, 0): P["svc"], (3, 0): P["mgmt"], (4, 0): P["pwr"]}
    for r in range(1, 5):
        fills[(r, 5)] = "FFF1BF"
    colors = {(r, 0): P["white"] for r in range(1, 5)}
    table(sl, 0.5, 3.65, [0.6, 0.95, 0.9, 1.1, 1.0, 1.3], rows, row_h=0.3, name="표-포트 연결", fills=fills, colors=colors,
          align=["c", "l", "l", "l", "l", "l"], bold_cols=(0,))
    cap(sl, 7.0, 3.35, "점검 결과표 (상태 색)")
    rows = [["점검 항목", "기준", "결과", "상태"],
            ["전원 이중화", "PSU 2/2 정상", "정상", "● 정상"],
            ["디스크", "장애 0", "1개 예측 장애", "● 주의"],
            ["팬 · 온도", "입구 27℃ 이하", "24℃", "● 정상"],
            ["펌웨어", "권고 버전", "미적용", "● 조치"]]
    colors = {(1, 3): P["ok"], (2, 3): P["warn"], (3, 3): P["ok"], (4, 3): P["crit"]}
    table(sl, 7.0, 3.65, [1.5, 1.6, 1.4, 1.3], rows, row_h=0.3, name="표-점검 결과", colors=colors, bold_cols=(3,))
    text(sl, 0.5, 5.45, 12.3, 1.2, [
        [("표 다루기  ", {"bold": True, "color": P["svc"]}),
         ("행 추가 = 마지막 칸에서 Tab · 열 너비 = 경계선 끌기 · 머리글 색 = [표 디자인 > 음영] · 줄무늬는 짝수 행 연회색(#F6F8FA)", {})],
        [("Excel에서 가져오기  ", {"bold": True, "color": P["svc"]}),
         ("Excel 범위 복사 → PPT 표 칸 선택 후 붙여넣기(서식 유지 X, 값만) → 이 표의 서식이 그대로 유지됩니다", {})]],
         9.5, False, P["s700"], line_spacing=1.2, spacing=4)
    return sl


# ── 비교 ────────────────────────────────────────────────────
def compare(prs):
    sl = slide(prs, "비교 — 변경 전·후, 안 비교", "구성 요소 · 비교")
    for k, (ttl, kind, tg) in enumerate((("현재 (As-Is)", "mute", "현재"), ("변경 후 (To-Be)", "info", "변경 후"))):
        x = 0.5 + k * 3.55
        rrect(sl, x, 1.35, 3.25, 3.3, 0.08, fill=P["s50"] if k == 0 else P["t_info"], line=None,
              name=f"비교 패널-{tg}")
        text(sl, x + 0.2, 1.48, 2.5, 0.3, ttl, 12, True, P["ink"] if k == 0 else P["svc"])
        for j in range(3):
            node_card(sl, x + 0.3, 1.95 + j * 0.62, "server", f"WEB-0{j+1}", "10G 단일" if k == 0 else "25G 이중화",
                      P["s500"] if k == 0 else P["svc"], w=2.65, h=0.5)
        text(sl, x + 0.2, 3.9, 2.9, 0.7, ["· 단일 스위치 연결" if k == 0 else "· TOR 이중화 (vPC)",
                                          "· 장애 시 서비스 중단" if k == 0 else "· 스위치 1대 장애에도 유지"],
             9, False, P["s700"], line_spacing=1.15)
    shape(sl, MSO_SHAPE.RIGHT_ARROW, 3.82, 2.75, 0.4, 0.45, fill=P["svc"], line=None, name="화살표")
    cap(sl, 7.9, 1.3, "안 비교 (권장안 강조)")
    rows = [["기준", "1안 증설", "2안 교체", "3안 클라우드"],
            ["비용", "●", "◐", "○"], ["성능", "◐", "●", "●"], ["작업 영향", "●", "◐", "◐"],
            ["운영 난이도", "●", "●", "○"], ["총평", "단기", "권장", "장기 검토"]]
    fills = {(r, 2): P["t_info"] for r in range(1, 6)}
    colors = {(5, 2): P["svc"]}
    table(sl, 7.9, 1.6, [1.3, 1.15, 1.15, 1.3], rows, row_h=0.36, name="표-안 비교", align=["l", "c", "c", "c"],
          fills=fills, colors=colors, zebra=False, size=10.5)
    tag(sl, 10.55, 1.38, "권장", "info", solid=True, w=0.5, h=0.22, size=8)
    text(sl, 7.9, 3.85, 5, 0.25, "● 좋음   ◐ 보통   ○ 미흡", 8.5, False, P["s500"])
    cap(sl, 0.5, 4.95, "장단점 (2단)")
    for k, (ttl, col, t, items) in enumerate((("장점", P["ok"], P["t_ok"], ["서비스 중단 없이 증설", "기존 운영 방식 유지", "도입 기간 2주"]),
                                              ("단점", P["crit"], P["t_crit"], ["랙 공간 4U 추가 필요", "전원 용량 재검토", "3년 후 재교체"]))):
        x = 0.5 + k * 3.55
        g = [rrect(sl, x, 5.25, 3.25, 1.45, 0.06, fill=t, line=None, name="바탕"),
             text(sl, x + 0.2, 5.35, 2, 0.25, ttl, 11, True, col)]
        for j, it in enumerate(items):
            g.append(text(sl, x + 0.2, 5.7 + j * 0.3, 2.9, 0.25, "· " + it, 9.5, False, P["ink"]))
        group(sl, g, f"장단점-{ttl}")
    return sl


# ── 현황 ────────────────────────────────────────────────────
def dashboard(prs):
    sl = slide(prs, "현황 — 지표 · 사용률 · 상태 · 추이", "구성 요소 · 현황")
    for k, (lab, val, unit, delta, good) in enumerate((("서비스 가용률", "99.98", "%", "▲ 0.02%p 전월 대비", True),
                                                       ("장애", "1", "건", "▼ 2건 전월 대비", True),
                                                       ("변경 작업", "14", "건", "계획 16 · 완료 14", True),
                                                       ("평균 CPU", "38", "%", "▲ 5%p 증가 추세", False))):
        kpi(sl, 0.5 + k * 3.1, 1.35, lab, val, unit, delta, good)
    cap(sl, 0.5, 2.75, "자원 사용률 — 막대 길이 · 색은 직접 조정 (70% 주의 · 85% 위험)")
    for k, (lab, pct) in enumerate((("CPU", 38), ("메모리", 72), ("스토리지", 86), ("네트워크 대역", 41), ("랙 전원", 64))):
        usage_bar(sl, 0.5, 3.1 + k * 0.36, lab, pct, w=6.0)
    cap(sl, 7.0, 2.75, "시스템 상태 (점 = 상태 색)")
    rows = [["시스템", "서버", "네트워크", "스토리지", "백업"],
            ["인터넷뱅킹", "●", "●", "●", "●"], ["내부 업무", "●", "●", "●", "●"], ["정보계", "●", "●", "●", "●"],
            ["DR", "●", "●", "●", "●"]]
    st = {(1, 1): P["ok"], (1, 2): P["ok"], (1, 3): P["ok"], (1, 4): P["ok"], (2, 1): P["ok"], (2, 2): P["warn"],
          (2, 3): P["ok"], (2, 4): P["ok"], (3, 1): P["ok"], (3, 2): P["ok"], (3, 3): P["crit"], (3, 4): P["ok"],
          (4, 1): P["ok"], (4, 2): P["ok"], (4, 3): P["ok"], (4, 4): P["s300"]}
    table(sl, 7.0, 3.08, [1.5, 1.05, 1.05, 1.05, 1.05], rows, row_h=0.3, name="표-상태", align=["l", "c", "c", "c", "c"],
          colors=st, size=11)
    text(sl, 7.0, 4.65, 5.8, 0.2, "● 정상   ● 주의   ● 장애   ● 해당 없음", 8.5, False, P["s500"])
    cap(sl, 0.5, 5.0, "추이 차트 (차트 우클릭 > 데이터 편집으로 값 수정)")
    cd = CategoryChartData()
    cd.categories = ["4월", "5월", "6월", "7월", "8월", "9월"]
    cd.add_series("장애", (3, 2, 4, 1, 3, 1))
    cd.add_series("변경", (11, 9, 15, 12, 16, 14))
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(0.5), I(5.25), I(6.2), I(1.65), cd)
    gf.name = "차트-월별 장애·변경"
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.RIGHT
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(8.5)
    ch.legend.font.name = "맑은 고딕"
    for s, col in zip(ch.series, (P["crit"], P["svc"])):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(col)
    pl = ch.plots[0]
    pl.gap_width = 80
    pl.has_data_labels = True
    pl.data_labels.font.size = Pt(8)
    pl.data_labels.font.name = "맑은 고딕"
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(8.5)
    ca.tick_labels.font.name = "맑은 고딕"
    ca.format.line.color.rgb = rgb(P["s300"])
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    text(sl, 7.0, 5.1, 5.8, 1.6, [
        [("요약 문장 예시", {"bold": True, "color": P["svc"]})],
        "· 9월 가용률 99.98% — 목표(99.95%) 달성",
        "· 스토리지 사용률 86% — 증설 검토 필요 (10월 안건)",
        "· 장애 1건: 9/14 SAN-B 포트 오류 (조치 완료, 영향 없음)"], 9.5, False, P["s700"], line_spacing=1.2)
    return sl


# ── 작업 절차 ────────────────────────────────────────────────
def procedure(prs):
    sl = slide(prs, "작업 절차 — 단계 · 시간표 · 체크리스트 · 롤백", "구성 요소 · 절차")
    steps(sl, 0.5, 1.35, [("사전 점검", "백업 확인\n모니터링 알람 중지"), ("작업 준비", "케이블 라벨 확인\n예비 부품 준비"),
                          ("작업 수행", "케이블 이설\n설정 변경"), ("정상 확인", "서비스 · 링크 확인\n로그 점검"),
                          ("종료 보고", "알람 원복\n결과 공유")])
    cap(sl, 0.5, 2.85, "시간표 (타임라인)")
    y = 3.55
    line(sl, 0.7, y, 7.3, y, P["s300"], 2.0, name="시간축")
    for k, (t, d, st) in enumerate((("22:00", "작업 시작 · 공지", "ok"), ("22:10", "사전 점검", "ok"), ("22:30", "케이블 이설", "info"),
                                    ("23:30", "정상 확인", "mute"), ("00:00", "종료 보고", "mute"))):
        x = 0.8 + k * 1.6
        col = {"ok": P["ok"], "info": P["svc"], "mute": P["s400"]}[st]
        oval(sl, x - 0.09, y - 0.09, 0.18, 0.18, fill=col, line=P["white"], lw=1.5, name=f"시점-{t}")
        text(sl, x - 0.6, y - 0.42, 1.2, 0.22, t, 10, True, P["ink"], align="c")
        text(sl, x - 0.7, y + 0.16, 1.4, 0.4, d, 8.5, False, P["s700"], align="c")
    text(sl, 0.7, 4.2, 6.5, 0.2, "● 완료   ● 진행 중   ● 예정", 8.5, False, P["s500"])
    cap(sl, 0.5, 4.65, "체크리스트")
    for k, (t, c) in enumerate((("변경 요청서 승인 (CR-2026-0931)", True), ("백업 완료 확인", True), ("관련 부서 공지", True),
                                ("롤백 절차 리허설", False), ("작업 후 모니터링 30분", False))):
        checkbox(sl, 0.5 + (k // 3) * 3.4, 4.98 + (k % 3) * 0.34, t, c, w=3.3)
    cap(sl, 7.9, 2.85, "판단 · 롤백 흐름")
    a = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 7.9, 3.2, 1.6, 0.5, fill=P["navy"], line=None, txt="작업 수행", size=10,
              bold=True, color=P["white"], name="흐름-작업")
    b = shape(sl, MSO_SHAPE.DIAMOND, 10.0, 3.05, 1.5, 0.8, fill=P["white"], line=P["navy"], lw=1.25, txt="정상?",
              size=10, bold=True, color=P["navy"], name="흐름-판단")
    c_ = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 11.9, 3.2, 1.0, 0.5, fill=P["ok"], line=None, txt="완료", size=10,
               bold=True, color=P["white"], name="흐름-완료")
    d_ = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 10.0, 4.55, 1.5, 0.5, fill=P["crit"], line=None, txt="롤백", size=10,
               bold=True, color=P["white"], name="흐름-롤백")
    connect(sl, a, 3, b, 1, "straight", P["s600"], 1.25, tail="triangle")
    connect(sl, b, 3, c_, 1, "straight", P["s600"], 1.25, tail="triangle")
    connect(sl, b, 2, d_, 0, "straight", P["s600"], 1.25, tail="triangle")
    text(sl, 11.5, 3.12, 0.4, 0.2, "예", 8.5, True, P["ok"])
    text(sl, 10.8, 4.0, 0.6, 0.2, "아니오", 8.5, True, P["crit"])
    text(sl, 7.9, 5.35, 5, 1.3, [
        [("롤백 기준 예", {"bold": True, "color": P["crit"]})],
        "· 링크 미기동 10분 이상 · 서비스 응답 없음",
        "· 롤백 시간 20분 이내 (원래 포트로 케이블 복구)"], 9.5, False, P["s700"], line_spacing=1.2)
    return sl
