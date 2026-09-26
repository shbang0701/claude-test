# -*- coding: utf-8 -*-
"""[랙] 시트: 장비 목록을 입력하면 전면/후면 실장도가 수식 + 조건부 서식으로 자동으로 그려진다."""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Border, Side, Font
from openpyxl.worksheet.datavalidation import DataValidation
from xlkit import Canvas, P, side, fill, font, align, COL, add_cf
from xlparts import _tw

TAB_RACK = "1F3A5F"
FIRST, UMAX = 6, 42
LAST = FIRST + UMAX - 1          # 47
LROWS = 40
LLAST = FIRST + LROWS - 1        # 45
KIND_TINT = [("서버", "DCE8F6"), ("네트워크", "D5EEEE"), ("스토리지", "FBE4D3"), ("보안", "F8DCDA"),
             ("전원", "EFE3F5"), ("선반", "ECE7DC"), ("기타", "E4E9EF")]
SIDES = "양면,전면,후면,선반,0U"
# 목록 열 — 시스템 = 여러 박스가 한 시스템일 때 같은 이름(예: 본체 + I/O 드로어 = AIX-01)
LC = dict(U=(34, 2), 높이=(36, 2), 면=(38, 2), 장비명=(40, 5), 시스템=(45, 3), 모델=(48, 6), 구분=(54, 3),
          시트=(57, 6), 이동=(63, 1))
NOT_DRAWN = ("0U", "선반")        # 실장도에 그리지 않는 면 (선반 위 장비는 선반 한 줄로 표시)
HC = dict(u=65, fidx=66, fcnt=67, ridx=68, rcnt=69)


def _rng(key):
    c = COL(LC[key][0])
    return f"${c}${FIRST}:${c}${LLAST}"


def build_rack(wb, rack="DC1-R02", info=None, devices=(), title_note=None):
    cv = Canvas(wb, "랙", TAB_RACK)
    ws = cv.ws
    info = info or {}
    cv.h(1, 24)
    cv.text(1, 2, rack, sz=16, b=True)
    cv.text(1, 2 + int(_tw(rack, 16)) + 2, "랙 실장도 · 장비 목록", sz=11, col=P["s700"])
    cv.text(1, 63, "목록을 입력하면 실장도가 자동으로 그려집니다", sz=8, col=P["s500"], h="right")
    for c in range(2, 64):
        cv.add_border(1, c, bottom=side("medium", P["navy"]))
    cv.h(2, 16)
    spec = [("설치 위치", 4, 9), ("규격", 2, 9), ("전원", 2, 12), ("관리 담당", 4, 7), ("작성", 2, 7)]
    c = 2
    for lab, lw, vw in spec:
        cv.text(2, c, lab, sz=7.5, col=P["s500"])
        cv.text(2, c + lw, info.get(lab, ""), sz=8.5, b=True)
        c += lw + vw + 1
    cv.h(3, 5)
    # 구역 제목
    cv.h(4, 14)
    cv.text(4, 2, "전면", sz=9, b=True, col=P["navy"])
    cv.text(4, 5, "앞에서 본 모습", sz=7, col=P["s500"])
    cv.text(4, 18, "후면", sz=9, b=True, col=P["navy"])
    cv.text(4, 21, "뒤에서 본 모습 (좌우 반대)", sz=7, col=P["s500"])
    cv.text(4, 34, "장비 목록", sz=9, b=True, col=P["navy"])
    cv.text(4, 39, "← 여기에 입력 (U = 장비 맨 아래 U 번호)", sz=7, col=P["s500"])
    # 머리글
    cv.h(5, 14)
    for c in list(range(2, 17)) + list(range(18, 33)):
        cv.put(5, c, style="표 머리글")
    cv.c(5, 2).value = "U"
    cv.c(5, 2).alignment = align("center")
    cv.c(5, 3).value = "장비 (맨 위 U에 이름 표시)"
    cv.c(5, 19).value = "장비"
    cv.c(5, 32).value = "U"
    cv.c(5, 32).alignment = align("center")
    for key, (c0, w) in LC.items():
        for x in range(c0, c0 + w):
            cv.put(5, x, style="표 머리글")
        cv.c(5, c0).value = {"시트": "시트 이름", "이동": "▶"}.get(key, key)
        if key in ("U", "높이", "이동"):
            cv.c(5, c0).alignment = align("center")
        cv.merge(5, c0, 5, c0 + w - 1)
    # U 행
    LU, LH, LS, LN, LM, LK = _rng("U"), _rng("높이"), _rng("면"), _rng("장비명"), _rng("모델"), _rng("구분")
    hh = f"({LH}+({LH}=0))"
    dotted = side("dotted", P["s300"])
    for i in range(UMAX):
        r = FIRST + i
        u = UMAX - i
        cv.h(r, 10)
        for uc in (2, 32):
            cv.put(r, uc, u, f=font(6.5, False, P["s500"]), fl=fill(P["s100"]), al=align("center"))
        for uc in (16, 18):
            cv.c(r, uc).fill = fill(P["s100"])
        for c in list(range(3, 16)) + list(range(19, 32)):
            cv.add_border(r, c, bottom=dotted)
        U = f"$BM{r}"
        ws.cell(r, HC["u"]).value = f"=B{r}"
        base = f"ISNUMBER({LU})*({LU}<={U})*({LU}+{hh}-1>={U})" + "".join(f'*({LS}<>"{x}")' for x in NOT_DRAWN)
        ws.cell(r, HC["fcnt"]).value = f'=SUMPRODUCT({base}*({LS}<>"후면"))'
        ws.cell(r, HC["fidx"]).value = f'=SUMPRODUCT({base}*({LS}<>"후면")*(ROW({LU})-{FIRST - 1}))'
        ws.cell(r, HC["rcnt"]).value = f'=SUMPRODUCT({base}*({LS}<>"전면"))'
        ws.cell(r, HC["ridx"]).value = f'=SUMPRODUCT({base}*({LS}<>"전면")*(ROW({LU})-{FIRST - 1}))'
        for dc, ic, cc in ((3, "BN", "BO"), (19, "BP", "BQ")):
            idx = f"${ic}{r}"
            cnt = f"${cc}{r}"
            top = f"INDEX({LU},{idx})+INDEX({LH},{idx})+(INDEX({LH},{idx})=0)-1"
            label = f'INDEX({LN},{idx})&IF(INDEX({LM},{idx})="",""," · "&INDEX({LM},{idx}))'
            cv.put(r, dc, f'=IF({cnt}<>1,IF({cnt}>1,"⚠ U 겹침",""),IF({U}={top},{label},""))',
                   f=font(7, True, P["ink"]), al=align("left", indent=1))
    # 외곽선(랙 레일) — 조건부 서식이 닿지 않는 레일 열과 위·아래 행에만 선을 둔다
    med = side("medium", P["s700"])
    for r in range(FIRST, LAST + 1):
        cv.add_border(r, 2, left=med)
        cv.add_border(r, 16, right=med)
        cv.add_border(r, 18, left=med)
        cv.add_border(r, 32, right=med)
    for c in list(range(2, 17)) + list(range(18, 33)):
        cv.add_border(FIRST - 1, c, bottom=med)
        cv.add_border(LAST + 1, c, top=med)
    # 조건부 서식 (앞/뒤) — 셀마다 규칙 하나만 맞도록 구성 (Excel·LibreOffice 동일하게 보임)
    thin = Side(style="thin", color=P["s700"])
    for c1, c2, ic, cc in ((3, 15, "BN", "BO"), (19, 31, "BP", "BQ")):
        rng = f"{COL(c1)}{FIRST}:{COL(c2)}{LAST}"
        idx = f"MAX(1,${ic}{FIRST})"
        cnt = f"${cc}{FIRST}"
        U = f"$BM{FIRST}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{cnt}>1"], stopIfTrue=True,
                                      fill=PatternFill(start_color=P["warn_bg"], end_color=P["warn_bg"], fill_type="solid"),
                                      font=Font(color=P["warn_fg"], bold=True)))
        top = f"INDEX({LU},{idx})+INDEX({LH},{idx})+(INDEX({LH},{idx})=0)-1"
        kinds = KIND_TINT + [("*", P["s150"])]
        for k, tint in kinds:
            cond = f'INDEX({LK},{idx})="{k}"' if k != "*" else "TRUE"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND({cnt}=1,{U}={top},{cond})"], stopIfTrue=True, border=Border(top=thin),
                fill=PatternFill(start_color=tint, end_color=tint, fill_type="solid")))
        for k, tint in kinds:
            cond = f'INDEX({LK},{idx})="{k}"' if k != "*" else "TRUE"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND({cnt}=1,{cond})"], stopIfTrue=True,
                fill=PatternFill(start_color=tint, end_color=tint, fill_type="solid")))
    # 장비 목록
    for i in range(LROWS):
        r = FIRST + i
        d = devices[i] if i < len(devices) else {}
        for key, (c0, w) in LC.items():
            for x in range(c0, c0 + w):
                cv.put(r, x, style="표 본문")
                cv.c(r, x).font = font(7.5, key == "장비명", P["ink"])
            v = d.get(key)
            if key == "이동":
                v = f'=IF({COL(LC["시트"][0])}{r}="","",HYPERLINK("#\'"&{COL(LC["시트"][0])}{r}&"\'!A1","▶"))'
                cv.c(r, c0).font = font(8, True, P["svc"])
            if v not in (None, ""):
                cv.c(r, c0).value = v
            if key in ("U", "높이", "이동"):
                cv.c(r, c0).alignment = align("center")
            if w > 1:
                cv.merge(r, c0, r, c0 + w - 1)
    # 목록 선택 + 구분 색
    for key, lst in (("면", SIDES), ("구분", ",".join(k for k, _ in KIND_TINT))):
        c0 = LC[key][0]
        dv = DataValidation(type="list", formula1=f'"{lst}"', allow_blank=True, showErrorMessage=False)
        dv.add(f"{COL(c0)}{FIRST}:{COL(c0)}{LLAST}")
        ws.add_data_validation(dv)
    kc = COL(LC["구분"][0])
    for k, tint in KIND_TINT:
        add_cf(ws, f"{kc}{FIRST}:{COL(LC['구분'][0] + 2)}{LLAST}", f'${kc}{FIRST}="{k}"', fill_color=tint)
    # U 겹침 표시(목록)
    # 요약·범례
    r = LAST + 2
    cv.h(r - 1, 6)
    cv.h(r, 14)
    cv.text(r, 2, "전면 사용", sz=7.5, col=P["s500"])
    cv.put(r, 6, f'=SUMPRODUCT(ISNUMBER({LU})*({LS}<>"후면")*({LS}<>"0U")*({LS}<>"선반")*{hh})&" U / {UMAX} U"',
           f=font(8.5, True, P["ink"]))
    cv.text(r, 13, "장비", sz=7.5, col=P["s500"])
    cv.put(r, 15, f'=(COUNTA({LN})-COUNTIF({LK},"선반"))&" 대"', f=font(8.5, True, P["ink"]))
    x = 21
    for k, tint in KIND_TINT:
        cv.c(r, x).fill = fill(tint)
        cv.c(r, x).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        cv.text(r, x + 1, k, sz=7, col=P["s700"])
        x += 2 + int(_tw(k, 7) + 1)
    cv.h(r + 1, 13)
    cv.text(r + 1, 2, "면: 양면 = 앞뒤 모두 차지 · 전면/후면 = 한쪽만 차지(짧은 장비) · 선반 = 선반·트레이 위 장비"
            "(선반은 구분 '선반'으로 한 줄 따로) · 0U = 세로 PDU  ·  U = 맨 아래 U · 높이 비우면 1U  ·  시스템 = 한 시스템인 박스끼리 같은 이름",
            sz=7, col=P["s500"])
    r += 1
    if title_note:
        cv.h(r + 1, 13)
        cv.text(r + 1, 2, title_note, sz=7, col=P["s500"])
        r += 1
    # 보조 열 숨김
    for c in HC.values():
        ws.column_dimensions[COL(c)].hidden = True
    cv.text(4, 65, "보조 계산 열 (숨김) — 지우지 마세요", sz=7, col=P["s500"])
    cv.finalize(extra_rows=10)
    cv.page_setup(63, r, title_rows="1:5", breaks=[])
    return cv
