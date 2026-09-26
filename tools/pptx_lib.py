# -*- coding: utf-8 -*-
"""PowerPoint 도형 조립 도우미 — 모든 요소는 편집 가능한 도형·텍스트로 만든다."""
import os, sys
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import design as D

W, H = 33.87, 19.05          # 슬라이드 크기(cm, 16:9)
M = 1.6                      # 좌우 여백
CW = W - 2 * M               # 본문 폭
TOP = 3.05                   # 본문 시작 y
BOT = 17.6                   # 본문 끝 y

def rgb(h): return RGBColor.from_string(h)

# ── 글꼴 ───────────────────────────────────────────────────────────────────
def _apply_font(run, size, bold, color, mono=False, italic=False, spc=None):
    f = run.font
    f.name = D.FONT_MONO if mono else D.FONT_LATIN
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        (latin.addnext(ea) if latin is not None else rPr.append(ea))
    ea.set('typeface', D.FONT_MONO_EA if mono else D.FONT_EA)
    if spc is not None: rPr.set('spc', str(int(spc * 100)))

def set_text(shape, text, size=11, bold=False, color=D.INK, align="l", anchor="ctr",
             mono=False, line=1.25, space_after=0, margins=(0.12, 0.12, 0.06, 0.06), wrap=True):
    """shape 의 텍스트를 한 가지 서식으로 채운다. 줄바꿈은 \\n."""
    tf = shape.text_frame
    tf.word_wrap = wrap
    l, r, t, bm = margins
    tf.margin_left, tf.margin_right = Cm(l), Cm(r)
    tf.margin_top, tf.margin_bottom = Cm(t), Cm(bm)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    lines = str(text).split("\n")
    tf.clear()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line
        if space_after: p.space_after = Pt(space_after)
        _apply_font(p.add_run(), size, bold, color, mono)
        p.runs[0].text = ln
    return shape

def rich(shape, parts, size=11, color=D.INK, align="l", anchor="ctr", line=1.25,
         margins=(0.12, 0.12, 0.06, 0.06), space_after=0):
    """parts: [(text, {size,bold,color,mono}), ...] — 한 문단 안에서 서식을 섞는다."""
    tf = shape.text_frame
    tf.word_wrap = True
    l, r, t, bm = margins
    tf.margin_left, tf.margin_right = Cm(l), Cm(r)
    tf.margin_top, tf.margin_bottom = Cm(t), Cm(bm)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    p.line_spacing = line
    if space_after: p.space_after = Pt(space_after)
    for txt, kw in parts:
        r_ = p.add_run(); r_.text = txt
        _apply_font(r_, kw.get("size", size), kw.get("bold", False),
                    kw.get("color", color), kw.get("mono", False))
    return shape

# ── 기본 도형 ──────────────────────────────────────────────────────────────
def _nostyle(sh):
    """테마 기본 효과(그림자·그라데이션)를 따르지 않도록 p:style 을 제거한다."""
    el = sh._element
    st = el.find(qn('p:style'))
    if st is not None: el.remove(st)

def _style(sh, fill, line, lw, dash=None):
    sh.shadow.inherit = False
    _nostyle(sh)
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln()
            for e in ln.findall(qn('a:prstDash')): ln.remove(e)
            d = ln.makeelement(qn('a:prstDash'), {'val': dash})
            ln.append(d)
    return sh

def box(s, x, y, w, h, fill=D.WHITE, line=D.LINE, lw=1.0, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        dash=None):
    sh = s.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sh.adjustments[0] = radius if radius is not None else min(0.5, 0.10 * min(w, h) / max(w, h) + 0.04)
        except Exception: pass
    _style(sh, fill, line, lw, dash)
    sh.text_frame.word_wrap = True
    return sh

def rect(s, x, y, w, h, fill=D.WHITE, line=D.LINE, lw=1.0, dash=None):
    return box(s, x, y, w, h, fill, line, lw, shape=MSO_SHAPE.RECTANGLE, dash=dash)

def tbox(s, x, y, w, h, text="", **kw):
    sh = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    sh.text_frame.word_wrap = True
    if text != "": set_text(sh, text, **kw)
    return sh

def line_h(s, x, y, w, color=D.LINE, lw=1.0, dash=None):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x), Cm(y), Cm(x + w), Cm(y))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(lw)
    ln.shadow.inherit = False; _nostyle(ln)
    if dash:
        l = ln.line._get_or_add_ln()
        l.append(l.makeelement(qn('a:prstDash'), {'val': dash}))
    return ln

def line_v(s, x, y, h, color=D.LINE, lw=1.0, dash=None):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x), Cm(y), Cm(x), Cm(y + h))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(lw)
    ln.shadow.inherit = False; _nostyle(ln)
    if dash:
        l = ln.line._get_or_add_ln()
        l.append(l.makeelement(qn('a:prstDash'), {'val': dash}))
    return ln

def arrow(s, x1, y1, x2, y2, color=D.INK_FAINT, lw=1.25, dash=None, head=True, tail=False,
          kind=MSO_CONNECTOR.STRAIGHT):
    c = s.shapes.add_connector(kind, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    c.shadow.inherit = False; _nostyle(c)
    ln = c.line._get_or_add_ln()
    if dash: ln.append(ln.makeelement(qn('a:prstDash'), {'val': dash}))
    if head: ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    if tail: ln.append(ln.makeelement(qn('a:headEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c

# ── 조합 부품 ──────────────────────────────────────────────────────────────
def node(s, x, y, w, h, title, sub=None, tag=None, accent=D.PRIMARY_MID, fill=D.WHITE,
         dash=None, tsize=11, ssize=8.5):
    """장비/서비스 노드: 제목 + 보조설명 + 왼쪽 색막대."""
    sh = rect(s, x, y, w, h, fill, D.LINE, 1.0, dash=dash)
    bar = rect(s, x, y, 0.13, h, accent, None)
    bar.line.fill.background()
    if sub:
        rich(sh, [(title + "\n", {"size": tsize, "bold": True, "color": D.INK}),
                  (sub, {"size": ssize, "color": D.INK_SOFT})], align="c", anchor="ctr", line=1.25,
             margins=(0.25, 0.15, 0.08, 0.08))
    else:
        set_text(sh, title, size=tsize, bold=True, color=D.INK, align="c",
                 margins=(0.25, 0.15, 0.08, 0.08))
    if tag:
        chip(s, x + w - 1.35, y - 0.28, 1.25, 0.56, tag, accent)
    return sh

def zone(s, x, y, w, h, label, color=D.PRIMARY_MID, fill=None, label_pos="t"):
    sh = box(s, x, y, w, h, fill, color, 1.0, radius=0.03, dash="dash")
    sh.text_frame.word_wrap = True
    if label_pos == "l":     # 왼쪽 가운데 (가로 밴드형)
        lb = tbox(s, x + 0.45, y, 3.0, h, label, size=10.5, bold=True, color=color, anchor="ctr",
                  margins=(0, 0, 0, 0))
    else:
        lb = tbox(s, x + 0.25, y + 0.12, 5.0, 0.6, label, size=8.5, bold=True, color=color, anchor="t",
                  margins=(0, 0, 0, 0))
    return sh, lb

def chip(s, x, y, w, h, text, color=D.PRIMARY_MID, fill=None, size=8.5):
    sh = box(s, x, y, w, h, fill or D.WHITE, color, 0.75, radius=0.5)
    set_text(sh, text, size=size, bold=True, color=color, align="c", margins=(0.05, 0.05, 0.02, 0.02))
    return sh

def pill(s, x, y, w, h, text, color, bg, size=9):
    sh = box(s, x, y, w, h, bg, None, 0, radius=0.5)
    set_text(sh, text, size=size, bold=True, color=color, align="c", margins=(0.08, 0.08, 0.02, 0.02))
    return sh

def badge(s, x, y, d, text, fill=D.PRIMARY_MID, color=D.WHITE, size=10):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(y), Cm(d), Cm(d))
    _style(sh, fill, None, 0)
    set_text(sh, text, size=size, bold=True, color=color, align="c", margins=(0, 0, 0, 0))
    return sh

def dot(s, x, y, d, color):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(y), Cm(d), Cm(d))
    return _style(sh, color, None, 0)

def card(s, x, y, w, h, title, lines=None, accent=D.PRIMARY, tint=None, tsize=11.5):
    sh = box(s, x, y, w, h, tint or D.WHITE, D.LINE, 1.0)
    top = rect(s, x, y, w, 0.12, accent, None); top.line.fill.background()
    t = tbox(s, x + 0.45, y + 0.42, w - 0.9, 0.8, title, size=tsize, bold=True, color=D.INK, anchor="t",
             margins=(0, 0, 0, 0))
    if lines:
        tb = tbox(s, x + 0.45, y + 1.35, w - 0.9, h - 1.7, "", anchor="t")
        tf = tb.text_frame; tf.word_wrap = True; tf.clear()
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.3; p.space_after = Pt(5)
            _apply_font(p.add_run(), 9.5, False, D.INK_SOFT)
            p.runs[0].text = ln
    return sh

def kpi(s, x, y, w, h, label, value, unit="", color=D.PRIMARY, sub=None, vsize=17):
    sh = box(s, x, y, w, h, D.WHITE, D.LINE, 1.0)
    tbox(s, x + 0.4, y + 0.35, w - 0.8, 0.6, label, size=9, bold=False, color=D.INK_FAINT, anchor="t",
         margins=(0, 0, 0, 0))
    v = tbox(s, x + 0.4, y + 0.88, w - 0.8, 1.5, "", anchor="t", margins=(0, 0, 0, 0))
    rich(v, [(value, {"size": vsize, "bold": True, "color": color}), (" " + unit, {"size": 10, "color": D.INK_SOFT})],
         align="l", anchor="t", margins=(0, 0, 0, 0))
    if sub:
        tbox(s, x + 0.4, y + h - 0.78, w - 0.8, 0.6, sub, size=8.5, color=D.INK_SOFT, anchor="t",
             margins=(0, 0, 0, 0))
    return sh

# ── 슬라이드 뼈대 ──────────────────────────────────────────────────────────
def slide(prs, title, sub=None, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tbox(s, M, 1.05, CW - 8, 1.1, title, size=19, bold=True, color=D.PRIMARY, anchor="t",
         margins=(0, 0, 0, 0))
    if sub:
        tbox(s, M, 2.05, CW - 8, 0.7, sub, size=10, color=D.INK_SOFT, anchor="t", margins=(0, 0, 0, 0))
    line_h(s, M, 2.72, CW, D.LINE, 1.0)
    if note:
        tbox(s, M, BOT + 0.35, CW, 0.6, note, size=8.5, color=D.INK_FAINT, anchor="t",
             margins=(0, 0, 0, 0))
    return s

def set_theme_fonts(prs):
    """테마 글꼴을 맑은 고딕으로 바꿔 새로 만든 텍스트 상자도 같은 글꼴이 되게 한다."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    master = prs.slide_masters[0]
    theme = master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    import re
    def fix(m):
        kind = m.group(1)
        return (f'<a:{kind}><a:latin typeface="{D.FONT_LATIN}"/><a:ea typeface="{D.FONT_EA}"/>'
                f'<a:cs typeface=""/>')
    xml = re.sub(r'<a:(majorFont|minorFont)><a:latin[^/]*/><a:ea[^/]*/><a:cs[^/]*/>', fix, xml)
    theme._blob = xml.encode("utf-8")

# ── 표 ─────────────────────────────────────────────────────────────────────
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

def _cell_lines(cell, color=D.LINE_SOFT, w=0.75, edges="b"):
    tcPr = cell._tc.get_or_add_tcPr()
    order = [("l", 'a:lnL'), ("r", 'a:lnR'), ("t", 'a:lnT'), ("b", 'a:lnB')]
    for key, tag in order:
        for e in tcPr.findall(qn(tag)): tcPr.remove(e)
    idx = 0
    for key, tag in order:
        ln = tcPr.makeelement(qn(tag), {'w': str(int(w * 12700)), 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        if key in edges:
            f = ln.makeelement(qn('a:solidFill'), {})
            c = ln.makeelement(qn('a:srgbClr'), {'val': color}); f.append(c); ln.append(f)
        else:
            ln.append(ln.makeelement(qn('a:noFill'), {}))
        tcPr.insert(idx, ln); idx += 1

def ptable(s, x, y, widths, headers, rows, row_h=0.95, head_h=1.0, size=9.5,
           head_fill=D.PRIMARY_TINT, align=None):
    """편집 가능한 PowerPoint 표. widths: cm 목록."""
    n_r, n_c = len(rows) + 1, len(widths)
    shp = s.shapes.add_table(n_r, n_c, Cm(x), Cm(y), Cm(sum(widths)), Cm(head_h + row_h * len(rows)))
    tbl = shp.table
    tbl._tbl.tblPr.find(qn('a:tableStyleId')).text = NO_STYLE
    for i, w in enumerate(widths): tbl.columns[i].width = Cm(w)
    tbl.rows[0].height = Cm(head_h)
    for r in range(1, n_r): tbl.rows[r].height = Cm(row_h)
    data = [headers] + rows
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left, cell.margin_right = Cm(0.25), Cm(0.18)
            cell.margin_top, cell.margin_bottom = Cm(0.12), Cm(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill if r == 0 else D.WHITE)
            if r == 0: _cell_lines(cell, D.PRIMARY_MID, 1.0, "b")
            else:      _cell_lines(cell, D.LINE_SOFT, 0.75, "b")
            tf = cell.text_frame; tf.word_wrap = True; tf.clear()
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[
                (align or ["l"] * n_c)[c]]
            p.line_spacing = 1.2
            _apply_font(p.add_run(), size, r == 0, D.PRIMARY if r == 0 else D.INK)
            p.runs[0].text = str(val)
    return shp

# ── 차트(편집 가능한 PowerPoint 기본 차트) ─────────────────────────────────
def _chart_style(chart, size=9.5, legend=True, gridlines=True, value_fmt=None):
    from pptx.enum.chart import XL_LEGEND_POSITION
    chart.font.size = Pt(size)
    chart.font.name = D.FONT_LATIN
    chart.font.color.rgb = rgb(D.INK_SOFT)
    try: chart.has_title = False
    except Exception: pass
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    try:
        va = chart.value_axis
        va.has_major_gridlines = gridlines
        if gridlines:
            gl = va.major_gridlines.format.line
            gl.color.rgb = rgb(D.LINE_SOFT); gl.width = Pt(0.75)
        va.format.line.color.rgb = rgb(D.LINE)
        va.has_minor_gridlines = False
        if value_fmt: va.tick_labels.number_format = value_fmt; va.tick_labels.number_format_is_linked = False
    except Exception: pass
    try:
        ca = chart.category_axis
        ca.format.line.color.rgb = rgb(D.LINE)
        ca.has_major_gridlines = False
    except Exception: pass

def bar_chart(s, x, y, w, h, cats, series, colors=None, stacked=False, legend=True,
              gap=60, overlap=None):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    t = XL_CHART_TYPE.COLUMN_STACKED_100 if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = s.shapes.add_chart(t, Cm(x), Cm(y), Cm(w), Cm(h), cd)
    ch = gf.chart
    _chart_style(ch, legend=legend and len(series) > 1)
    plot = ch.plots[0]
    plot.gap_width = gap
    if overlap is not None: plot.overlap = overlap
    palette = colors or [D.PRIMARY_MID, D.PRIMARY, D.INK_FAINT, D.ST_OK, D.ST_WARN, D.ST_BAD]
    for i, sr in enumerate(ch.series):
        sr.format.fill.solid(); sr.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
        sr.format.line.fill.background()
    return gf

def line_chart(s, x, y, w, h, cats, series, colors=None, legend=True, markers=True):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    t = XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE
    gf = s.shapes.add_chart(t, Cm(x), Cm(y), Cm(w), Cm(h), cd)
    ch = gf.chart
    _chart_style(ch, legend=legend and len(series) > 1)
    palette = colors or [D.PRIMARY, D.PRIMARY_MID, D.INK_FAINT]
    for i, sr in enumerate(ch.series):
        sr.format.line.color.rgb = rgb(palette[i % len(palette)])
        sr.format.line.width = Pt(2.0)
        sr.smooth = False
    return gf

# ── 도형 도구상자용 추가 도우미 ────────────────────────────────────────────
def alpha(shape, color, pct):
    """반투명 채우기(사진 위 글자 상자 등). pct: 0~100 (불투명도)."""
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sf.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    clr.append(a)
    return shape

def shp(s, kind, x, y, w, h, fill=D.PRIMARY_TINT, line=D.PRIMARY_MID, lw=1.0, adj=None,
        text=None, size=9.5, color=None, bold=True, rot=None):
    sh = s.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    _style(sh, fill, line, lw)
    if adj is not None:
        try:
            for i, v in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
                sh.adjustments[i] = v
        except Exception: pass
    if rot: sh.rotation = rot
    if text is not None:
        set_text(sh, text, size=size, bold=bold, color=color or D.PRIMARY, align="c",
                 margins=(0.1, 0.1, 0.02, 0.02))
    return sh

HEADS = {"삼각": "triangle", "스텔스": "stealth", "열린": "arrow", "원": "oval", "마름모": "diamond"}

def conn(s, pts, color=D.PRIMARY_MID, lw=1.5, dash=None, head="triangle", tail=None,
         kind=MSO_CONNECTOR.STRAIGHT, hw="med", hl="med"):
    """(x1,y1,x2,y2) 연결선. kind 로 직선·꺾인선·곡선을 고른다."""
    x1, y1, x2, y2 = pts
    c = s.shapes.add_connector(kind, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    c.shadow.inherit = False; _nostyle(c)
    ln = c.line._get_or_add_ln()
    if dash: ln.append(ln.makeelement(qn('a:prstDash'), {'val': dash}))
    if tail: ln.append(ln.makeelement(qn('a:headEnd'), {'type': tail, 'w': hw, 'len': hl}))
    if head: ln.append(ln.makeelement(qn('a:tailEnd'), {'type': head, 'w': hw, 'len': hl}))
    return c

def label(s, x, y, w, text, size=8.5, color=D.INK_FAINT, align="l", bold=False):
    return tbox(s, x, y, w, 0.6, text, size=size, bold=bold, color=color, align=align,
                anchor="t", margins=(0, 0, 0, 0))
